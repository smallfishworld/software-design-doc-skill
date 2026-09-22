#!/usr/bin/env python3
"""Rebuild the bundled A4 Chinese HLD template. Requires python-docx.

Run only when maintaining this repository. Document generation must edit a copy
of the template. Pass --output to create a separate preview instead.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def field(paragraph, instruction, hint=""):
    for kind in ("begin", "instruction", "separate", "text", "end"):
        run = paragraph.add_run()
        if kind == "instruction":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instruction
        elif kind == "text":
            run.text = hint
            continue
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def table(doc, headers, rows, widths, *, alignment=WD_TABLE_ALIGNMENT.LEFT):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = alignment
    t.style = "Table Grid"

    width_twips = [Cm(width).twips for width in widths]
    tbl_w = t._tbl.tblPr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        t._tbl.tblPr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(width_twips)))

    for grid_col, twips in zip(t._tbl.tblGrid.gridCol_lst, width_twips):
        grid_col.set(qn("w:w"), str(twips))
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri, row in enumerate(t.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell, width, twips in zip(row.cells, widths, width_twips):
            cell.width = Cm(width)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(twips))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EDF2")
                cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(5)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            el.set(qn(f"w:{attr}"), value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def prompt(doc, text):
    p = doc.add_paragraph(text)
    for r in p.runs:
        r.font.color.rgb = RGBColor.from_string("595959")
    return p


def build(output: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1)
    for name, size in (("Normal", 11), ("Title", 26), ("Subtitle", 14),
                       ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
        st.paragraph_format.space_after = Pt(6)
        if name.startswith("Heading"):
            st.paragraph_format.space_before = Pt(10)
            st.paragraph_format.keep_with_next = True
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    # Some python-docx distributions ship a styled default with title borders.
    # Explicitly clear those so rebuilding does not inherit decorative rules.
    for st in doc.styles:
        for border in list(st.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Subtitle"].font.italic = False
    header = sec.header.paragraphs[0]
    header.text = "{{PROJECT}}  软件概要设计说明书  {{VER}}"
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, " PAGE ", "1")
    sec.different_first_page_header_footer = True
    doc.core_properties.title = "软件概要设计说明书模板"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 9, 16, tzinfo=timezone.utc)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)

    p = doc.add_paragraph("{{PROJECT}}", "Subtitle")
    p.paragraph_format.space_before = Pt(80)
    doc.add_paragraph("软件概要设计说明书", "Title")
    doc.add_paragraph("Software High-Level Design Specification", "Subtitle")
    doc.add_paragraph("系统边界、模块职责、接口与关键设计决策", "Normal")
    table(doc, ["文档信息", "内容"], [
        ["文档编号", "{{DOC_ID}}"], ["版本", "{{VER}}"], ["编制", "{{AUTHOR}}"],
        ["审核与批准", "{{REVIEWER}} / {{APPROVER}}"], ["日期", "{{DATE}}"],
        ["密级与状态", "{{CONF}} / {{STATUS}}"],
    ], [3.2, 8.4], alignment=WD_TABLE_ALIGNMENT.CENTER)

    doc.add_page_break()
    doc.add_paragraph("文档控制", "Title")
    table(doc, ["属性", "内容"], [["适用范围", "{{SCOPE}}"], ["所属部门", "{{DEPT}}"]], [4, 12.6])
    doc.add_paragraph("修订记录", "Subtitle")
    table(doc, ["版本", "日期", "修订内容", "编制与审核"],
          [["{{VER}}", "{{DATE}}", "{{CHANGE}}", "{{AUTHOR}} / {{REVIEWER}}"]], [2, 3, 7, 4.6])
    doc.add_paragraph("目录", "Subtitle")
    field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', "在 Word 中更新域以生成目录")
    prompt(doc, "模板使用说明：替换灰色提示与示例行；未知元数据标为待确认。第 8～11 章为条件候选章节：全局影响时保留，局部影响时融合，无关时删除；保留时去掉标题中的‘条件候选’，并更新编号、交叉引用与目录。")

    doc.add_page_break()
    doc.add_heading("1 引言", 1)
    doc.add_heading("1.1 编写目的与范围", 2)
    prompt(doc, "说明文档用途、读者、系统边界及不在本次设计范围内的内容。")
    doc.add_heading("1.2 术语与缩略语", 2)
    table(doc, ["术语", "英文全称", "中文含义"], [["{{TERM}}", "{{FULL_NAME}}", "{{DESC}}"]], [3, 7, 6.6])
    doc.add_heading("1.3 参考资料", 2)
    table(doc, ["资料", "版本或位置", "用途"], [["{{REFERENCE}}", "{{LOCATION}}", "{{USE}}"]], [6, 5, 5.6])
    doc.add_heading("2 系统概述", 1)
    prompt(doc, "描述背景、用户、核心能力、外部系统、运行环境和关键约束；区分现状与目标。")
    doc.add_heading("3 设计目标与原则", 1)
    prompt(doc, "只保留影响本项目架构的目标，说明对应需求、约束及可验证的设计响应。")

    doc.add_page_break()
    doc.add_heading("4 软件总体架构", 1)
    doc.add_heading("4.1 架构与系统边界", 2)
    prompt(doc, "描述架构风格、系统边界及依赖方向；在此插入已验证的架构图与图题。")
    doc.add_heading("4.2 模块划分与职责", 2)
    table(doc, ["模块", "职责", "依赖与接口"], [["{{MODULE}}", "{{RESP}}", "{{DEP}}"]], [3.5, 6, 7.1])
    doc.add_heading("4.3 关键控制与数据流", 2)
    prompt(doc, "说明端到端关键路径、初始化、资源所有权和故障路径；必要时插入流程图或时序图。")
    doc.add_heading("4.4 关键设计决策", 2)
    prompt(doc, "说明约束、方案、备选项、选择理由和后果。并发、状态、实时性及部署内容影响全局架构时使用对应候选章节；仅局部影响时融入本章或模块章节。")
    doc.add_heading("5 模块概要设计", 1)
    doc.add_heading("5.1 {{MODULE}}", 2)
    prompt(doc, "按模块复制本节。说明职责、提供接口、依赖、输入输出、数据所有权、生命周期、异常及重要约束；避免逐函数描述。")
    doc.add_heading("6 数据设计", 1)
    table(doc, ["数据", "用途与结构", "所有者与生命周期"], [["{{DATA}}", "{{DESC}}", "{{OWNER}} / {{LIFE}}"]], [3.5, 6, 7.1])
    prompt(doc, "按需说明持久化、缓冲、缓存、一致性、序列化和兼容性。")

    doc.add_page_break()
    doc.add_heading("7 接口设计", 1)
    table(doc, ["接口", "提供方与使用方", "契约与失败语义"], [["{{IFACE}}", "{{PROVIDER}} / {{CONSUMER}}", "{{CONTRACT}}"]], [3.5, 5, 8.1])
    prompt(doc, "区分内部和外部接口，描述数据形态、同步性、所有权、错误、超时及版本兼容；未确认的数值标为设计建议或待确认。")
    doc.add_heading("8 多任务与并发设计（条件候选）", 1)
    prompt(doc, "跨模块的执行上下文、通信、资源所有权、同步、调度域、背压或停机策略影响总体架构时保留；局部内容移入对应模块/接口/数据章节，无关时删除。线程优先级、周期、栈大小、队列深度和锁顺序等实现参数通常属于详细设计。")
    doc.add_heading("9 状态行为设计（条件候选）", 1)
    prompt(doc, "统一生命周期、运行模式、协议/工作流或恢复状态协调多个模块或构成外部契约时保留；局部状态移入对应模块，无关时删除。完整转移矩阵与处理器动作通常属于详细设计。")
    doc.add_heading("10 性能、实时性与资源设计（条件候选）", 1)
    prompt(doc, "端到端时延、吞吐、截止期、容量或 CPU/内存/存储/网络预算驱动跨模块划分和契约时保留；局部约束移入对应模块/接口/数据章节，无关时删除。区分需求、建议预算、实测值和待确认项，不虚构数值。")
    doc.add_heading("11 部署、升级与兼容性设计（条件候选）", 1)
    prompt(doc, "节点/进程/容器/固件映射、制品划分、发布回滚、迁移、升级可用性、ABI、交叉编译或工具链约束影响架构时保留；局部约束移入总体架构、接口或风险章节，无关时删除。常规编译命令、打包和安装步骤不属于概要设计。")
    doc.add_heading("12 异常与故障处理", 1)
    table(doc, ["故障场景", "检测与影响", "恢复与降级"], [["{{SCENARIO}}", "{{IMPACT}}", "{{RECOVER}}"]], [4, 6, 6.6])
    doc.add_heading("13 日志与可观测性", 1)
    prompt(doc, "说明诊断、日志、指标和告警对关键路径的作用；根据项目规模裁剪。")
    doc.add_heading("14 安全设计", 1)
    prompt(doc, "可选：存在权限、敏感数据、网络暴露或完整性需求时保留，说明信任边界和具体措施。无关时删除并重新编号。")
    doc.add_heading("15 可测试性设计", 1)
    prompt(doc, "说明测试替身、故障注入、可观察状态、模拟及关键验证方法。")
    doc.add_heading("16 可维护性与扩展性", 1)
    prompt(doc, "说明变化点、扩展接口、配置策略、依赖隔离及兼容策略。")
    doc.add_heading("17 风险约束与待确认项", 1)
    table(doc, ["编号与类型", "问题与影响", "下一步"], [["{{ID}} / {{TYPE}}", "{{DESC}}", "{{ACTION}}"]], [3.5, 7, 6.1])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "templates/default-software-design-template.docx")
    args = parser.parse_args()
    build(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
